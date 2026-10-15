"""Explicit imports, part-code fidelity, stock planning, and atomic writes."""
from __future__ import annotations

import importlib.util
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import inventory
from inventory import (InventoryCancelled, InventoryError, InventoryStore,
                       inspect_spreadsheet, parse_excel, validate_rows)


class InventoryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.store = InventoryStore(self.root / "stock.sqlite", clock=lambda: 1_800_000_000)

    def tearDown(self):
        self.directory.cleanup()

    def csv(self, content):
        path = self.root / "parts.csv"
        path.write_text(content, encoding="utf-8-sig")
        return path

    def test_csv_plan_preserves_zero_codes_and_never_changes_stock(self):
        path = self.csv("Part Number,Description,Qty,Bin,Min Stock\n00123,Oil filter,2,Rack A,3\n")
        preview = inspect_spreadsheet(path)
        self.assertEqual(preview.headers[:2], ["Part Number", "Description"])
        self.assertEqual(preview.mapping["sku"], 0)
        plan = parse_excel(path)
        self.assertTrue(plan.valid)
        self.assertEqual(plan.rows[0]["sku"], "00123")
        self.assertEqual(self.store.summary()["item_count"], 0)
        changes = self.store.preview_import(plan)
        self.assertEqual(changes["created"], 1)
        self.assertEqual(self.store.summary()["revision"], 0)
        result = self.store.import_rows(plan.rows, plan.source, changes["revision"])
        self.assertEqual(result["created"], 1)
        self.assertEqual(self.store.get("00123")["location"], "Rack A")
        self.assertEqual(result["last_import_source"], "parts.csv")

    def test_header_mapping_can_be_selected_explicitly(self):
        path = self.csv("Code here,Number here,Where\n0007,4,C2\n")
        self.assertFalse(parse_excel(path).valid)
        plan = parse_excel(path, {"sku": "Code here", "quantity": "Number here", "location": "Where"})
        self.assertTrue(plan.valid)
        self.assertEqual(plan.rows[0], {"sku": "0007", "quantity": 4, "location": "C2"})

    def test_ambiguous_headers_require_explicit_index(self):
        path = self.csv("SKU,SKU,Qty\nA,B,2\n")
        self.assertFalse(parse_excel(path).valid)
        self.assertFalse(parse_excel(path, {"sku": "SKU", "quantity": "Qty"}).valid)
        self.assertTrue(parse_excel(path, {"sku": 1, "quantity": 2}).valid)

    def test_duplicate_codes_report_original_excel_row(self):
        plan = parse_excel(self.csv("SKU,Qty\nA,1\n\n a,3\n"))
        self.assertFalse(plan.valid)
        self.assertIn("Row 4", plan.errors[0])
        self.assertIn("row 2", plan.errors[0])
        with self.assertRaisesRegex(InventoryError, "duplicate part number"):
            self.store.import_rows([{"sku": "A", "quantity": 1}, {"sku": "a", "quantity": 2}])
        self.assertEqual(self.store.summary()["item_count"], 0)

    def test_invalid_counts_and_missing_quantities_reject_entire_import(self):
        for value in (True, -1, 1.5, "NaN", "Infinity", "2,2", 2**31):
            with self.subTest(value=value), self.assertRaises(InventoryError):
                self.store.import_rows([{"sku": "okay", "quantity": 2}, {"sku": "bad", "quantity": value}])
        for row in ({"sku": "bad"}, {"sku": "bad", "quantity": ""}, {"sku": "", "quantity": 2}):
            with self.assertRaises(InventoryError):
                self.store.import_rows([row])
        self.assertEqual(self.store.summary()["revision"], 0)
        self.assertEqual(self.store.summary()["item_count"], 0)

    def test_whole_counts_and_grouping(self):
        result = validate_rows([{"sku": "A", "quantity": "1,000.00", "target_stock": "1200"}])
        self.assertEqual(result[0]["quantity"], 1000)
        self.assertEqual(result[0]["target_stock"], 1200)

    def test_nested_values_are_rejected(self):
        for field in ("sku", "description", "location", "supplier"):
            with self.subTest(field=field), self.assertRaises(InventoryError):
                validate_rows([{"sku": "A", "quantity": 2, field: {"unexpected": "object"}}])

    def test_partial_upsert_preserves_location_reorder_and_supplier(self):
        self.store.import_rows([{"sku": "A", "quantity": 3, "location": "Shelf 7", "reorder_point": 5,
                                 "target_stock": 12, "reorder_quantity": 6, "supplier": "Acme"}])
        result = self.store.import_rows([{"sku": "A", "quantity": 2}])
        self.assertEqual(result["updated"], 1)
        item = self.store.get("a")
        self.assertEqual(item["location"], "Shelf 7")
        self.assertEqual(item["supplier"], "Acme")
        self.assertEqual(item["suggested_restock"], 10)
        self.assertEqual(len(self.store.low_stock()), 1)
        result = self.store.import_rows([{"sku": "A", "quantity": 2}])
        self.assertEqual(result["unchanged"], 1)

    def test_stale_preview_cannot_overwrite_new_stock(self):
        self.store.import_rows([{"sku": "A", "quantity": 3}])
        preview = self.store.preview_import([{"sku": "A", "quantity": 1}])
        self.store.import_rows([{"sku": "A", "quantity": 9}])
        with self.assertRaises(InventoryError) as caught:
            self.store.import_rows([{"sku": "A", "quantity": 1}], expected_revision=preview["revision"])
        self.assertEqual(caught.exception.status, 409)
        self.assertEqual(self.store.get("A")["quantity"], 9)

    def test_search_literal_wildcards_and_multiple_terms(self):
        self.store.import_rows([{"sku": "A_1", "description": "Oil filter", "location": "Rack 2", "quantity": 0},
                                {"sku": "AX1", "description": "Brake pad", "location": "Rack 1", "quantity": 2}])
        self.assertEqual([row["sku"] for row in self.store.search("A_1")], ["A_1"])
        self.assertEqual([row["sku"] for row in self.store.search("oil rack")], ["A_1"])
        self.assertEqual([row["sku"] for row in self.store.search(low_stock=True)], ["A_1"])
        self.assertEqual(self.store.search("%"), [])

    def test_reopen_is_durable_and_snapshot_replaces_stale_rows(self):
        self.store.import_rows([{"sku": "A", "quantity": 3}])
        reopened = InventoryStore(self.root / "stock.sqlite")
        self.assertEqual(reopened.get("A")["quantity"], 3)
        reopened.replace_snapshot([{"sku": "B", "quantity": 8}], revision=40,
                                  last_import_at="2026-10-14T12:00:00Z", last_import_source="office.xlsx")
        self.assertIsNone(reopened.get("A"))
        self.assertEqual(reopened.summary()["revision"], 40)
        reopened.replace_snapshot([], revision=41)
        self.assertEqual(reopened.summary()["item_count"], 0)

    def test_snapshot_preserves_record_provenance_and_rejects_bad_snapshot_atomically(self):
        self.store.replace_snapshot([{"sku": "A", "quantity": 3, "revision": 4,
                                      "updated_at": "2026-10-13T10:00:00Z", "source": "earlier.csv"}],
                                    revision=8, last_import_at="2026-10-14T12:00:00Z", last_import_source="latest.xlsx")
        item = self.store.get("A")
        self.assertEqual(item["revision"], 4)
        self.assertEqual(item["source"], "earlier.csv")
        self.assertEqual(item["updated_at"], "2026-10-13T10:00:00Z")
        with self.assertRaises(InventoryError):
            self.store.replace_snapshot([{"sku": "B", "quantity": 2, "revision": 99}], revision=9)
        self.assertIsNotNone(self.store.get("A"))
        self.assertIsNone(self.store.get("B"))

    def test_cumulative_item_capacity_is_atomic_but_updates_are_allowed(self):
        with patch.object(inventory, "MAX_ROWS", 2):
            self.store.import_rows([{"sku": "A", "quantity": 3}, {"sku": "B", "quantity": 4}])
            self.store.import_rows([{"sku": "A", "quantity": 7}])
            with self.assertRaises(InventoryError):
                self.store.import_rows([{"sku": "B", "quantity": 8}, {"sku": "C", "quantity": 2}])
        self.assertEqual(self.store.get("B")["quantity"], 4)

    def test_cancelled_and_oversize_files_do_not_read_stock(self):
        path = self.csv("SKU,Qty\nA,2\n")
        event = threading.Event()
        event.set()
        with self.assertRaises(InventoryCancelled):
            inspect_spreadsheet(path, cancel=event)
        with patch.object(inventory, "MAX_FILE_BYTES", 2), self.assertRaisesRegex(InventoryError, "50 MiB"):
            parse_excel(path)
        self.assertEqual(self.store.summary()["item_count"], 0)

    def test_csv_quoted_multiline_and_semicolon(self):
        plan = parse_excel(self.csv('SKU;Description;Qty\n001;"Two\nlines";3\n'))
        self.assertTrue(plan.valid)
        self.assertEqual(plan.rows[0]["description"], "Two\nlines")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "Install desktop spreadsheet requirements")
    def test_xlsx_formatted_codes_sheet_choice_and_formula_not_executed(self):
        import openpyxl
        book = openpyxl.Workbook()
        first = book.active
        first.title = "Empty"
        first.append(["SKU", "Qty"])
        selected = book.create_sheet("Parts")
        selected.append(["Part #", "Quantity", "Description"])
        selected.append([123, 4, "Safe values"])
        selected["A2"].number_format = "000000"
        selected.append(["FORMULA", "=1+1", "No cached value"])
        path = self.root / "parts.xlsx"
        book.save(path)
        preview = inspect_spreadsheet(path, sheet="Parts")
        self.assertEqual(preview.sheet_names, ["Empty", "Parts"])
        self.assertEqual(preview.sheet, "Parts")
        plan = parse_excel(path, sheet="Parts")
        self.assertFalse(plan.valid)
        self.assertEqual(plan.rows[0]["sku"], "000123")
        self.assertIn("Row 3", plan.errors[0])
        self.assertIn("quantity", plan.errors[0])
        with self.assertRaisesRegex(InventoryError, "worksheet"):
            parse_excel(path, sheet="Missing")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "Install desktop spreadsheet requirements")
    def test_xlsm_values_read_without_loading_vba(self):
        import openpyxl
        book = openpyxl.Workbook()
        book.active.append(["SKU", "Qty"])
        book.active.append(["001", 8])
        path = self.root / "parts.xlsm"
        book.save(path)
        self.assertEqual(parse_excel(path).rows[0]["sku"], "001")

    def test_network_database_path_is_rejected(self):
        with self.assertRaises(InventoryError):
            InventoryStore(r"\\server\shared\inventory.sqlite")


if __name__ == "__main__":
    unittest.main()
