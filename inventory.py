"""Local, structured parts inventory and explicit spreadsheet import planning.

Workbooks are read for cell values only. Import plans never change stock; callers
must explicitly commit them. This module has no AI, network, or order side effects.
"""
from __future__ import annotations

from contextlib import closing
import csv
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import io
import json
import math
from pathlib import Path
import re
import sqlite3
import time
import zipfile


MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 100
MAX_HISTORY_BYTES = 768 * 1024
FIELDS = ("sku", "description", "quantity", "location", "reorder_point",
          "reorder_quantity", "target_stock", "supplier", "barcode",
          "alternate_part_number", "substitute_part_number", "retail_price", "cost_price")
TEXT_LIMITS = {"sku": 128, "description": 2000, "location": 500, "supplier": 500, "barcode": 128,
               "alternate_part_number": 128, "substitute_part_number": 128}
COUNT_LIMIT = 2**31 - 1
ALIASES = {
    "sku": ("sku", "part", "partnumber", "partno", "partcode", "itemcode", "itemnumber", "itemno", "productcode", "stockcode"),
    "description": ("description", "item", "itemname", "product", "productname", "partname", "name"),
    "quantity": ("quantity", "qty", "stock", "onhand", "quantityonhand", "stockquantity", "instock", "balance"),
    "location": ("location", "bin", "binlocation", "shelf", "shelflocation", "aisle", "rack"),
    "reorder_point": ("reorderpoint", "reorderlevel", "minimum", "minimumstock", "minstock", "lowstock", "minqty"),
    "reorder_quantity": ("reorderquantity", "reorderqty", "orderquantity", "restockquantity", "restockqty"),
    "target_stock": ("targetstock", "targetquantity", "targetqty", "maximumstock", "maxstock"),
    "supplier": ("supplier", "vendor", "suppliername", "vendorname"),
    "alternate_part_number": ("alternatepartnumber", "alternatepartno", "altpartnumber", "altpartno", "alternativepartnumber"),
    "substitute_part_number": ("substitutepartnumber", "substitutepartno", "substitutepart", "replacementpartnumber", "replacementpartno"),
    "retail_price": ("retailprice", "retail", "sellingprice", "sellprice", "saleprice"),
    "cost_price": ("costprice", "cost", "purchaseprice", "buyprice", "unitcost"),
    "barcode": ("barcode", "barcodes", "upc", "upca", "ean", "ean13", "gtin", "gtin12", "gtin13", "gtin14"),
}


class InventoryError(ValueError):
    def __init__(self, detail, status=422):
        super().__init__(detail)
        self.detail, self.status = detail, status


class InventoryCancelled(InventoryError):
    pass


@dataclass
class InventoryPreview:
    headers: list[str]
    sample_rows: list[list]
    sheet_names: list[str]
    total_rows: int
    mapping: dict = field(default_factory=dict)
    source: str = ""
    sheet: str = ""


@dataclass
class InventoryImportPlan:
    rows: list[dict]
    errors: list[str]
    warnings: list[str]
    source: str
    headers: list[str]
    mapping: dict
    sheet: str = ""

    @property
    def valid(self):
        return bool(self.rows) and not self.errors


def _cancelled(cancel):
    if cancel is not None and (cancel() if callable(cancel) else cancel.is_set()):
        raise InventoryCancelled("Inventory import was cancelled.")


def _header_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower().replace("#", "number"))


def suggest_mapping(headers):
    result = {}
    for target, aliases in ALIASES.items():
        matches = [index for index, header in enumerate(headers) if _header_key(header) in aliases]
        if len(matches) == 1:
            result[target] = matches[0]
    return result


def _text(value, name, required=False):
    if value is None:
        value = ""
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise InventoryError(f"{name.replace('_', ' ').title()} must contain text or a simple number.")
    if not isinstance(value, str):
        value = str(value)
    value = value.strip()
    if (required and not value) or len(value) > TEXT_LIMITS[name] or any(ord(c) < 32 and c not in "\t\n\r" for c in value):
        raise InventoryError(f"{name.replace('_', ' ').title()} must contain {'1–' if required else 'at most '}{TEXT_LIMITS[name]} readable characters.")
    return value


def _count(value, name, nullable=False):
    if value is None or value == "":
        return None if nullable else 0
    if isinstance(value, bool):
        raise InventoryError(f"{name.replace('_', ' ').title()} must be a nonnegative whole number.")
    if isinstance(value, str):
        value = value.strip()
        if "," in value:
            if not re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.0+)?", value):
                raise InventoryError(f"{name.replace('_', ' ').title()} has an invalid number format.")
            value = value.replace(",", "")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise InventoryError(f"{name.replace('_', ' ').title()} must be a nonnegative whole number.") from None
    if not number.is_finite() or number < 0 or number > COUNT_LIMIT or number != number.to_integral_value():
        raise InventoryError(f"{name.replace('_', ' ').title()} must be a nonnegative whole number up to {COUNT_LIMIT}.")
    return int(number)


def _code(value, *, optional=False):
    """Scanner identifiers are text, never numbers or approximate searches."""
    if not isinstance(value, str):
        raise InventoryError("Barcode must be text so leading zeroes are preserved.")
    value = value.strip()
    if (not optional and not value) or len(value) > 128 or any(ord(c) < 32 or ord(c) == 127 for c in value):
        raise InventoryError("Enter a barcode or exact part number with 1–128 readable characters.")
    return value


def _price(value, name):
    label = name.replace('_', ' ').title()
    if value is None or isinstance(value, str) and not value.strip():
        return '0.00'
    if isinstance(value, bool):
        raise InventoryError(f'{label} must be a nonnegative amount with at most two decimal places.')
    text = str(value).strip()
    if text.startswith('$'):
        text = text[1:].strip()
    if ',' in text:
        if not re.fullmatch(r'\d{1,3}(?:,\d{3})+(?:\.\d+)?', text):
            raise InventoryError(f'{label} has an invalid number format.')
        text = text.replace(',', '')
    try:
        amount = Decimal(text)
        if (not amount.is_finite() or amount < 0 or amount > Decimal('999999999.99')
                or amount != amount.quantize(Decimal('0.01'))):
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        raise InventoryError(f'{label} must be a nonnegative amount with at most two decimal places.') from None
    return format((abs(amount) if amount == 0 else amount).quantize(Decimal('0.01')), '.2f')


def validate_row(row):
    if not isinstance(row, dict) or any(key not in FIELDS for key in row):
        raise InventoryError("Inventory rows must contain only recognized stock fields.")
    if "sku" not in row or "quantity" not in row or row["quantity"] is None or row["quantity"] == "":
        raise InventoryError("Each stock row needs a part number and quantity.")
    result = {"sku": _text(row["sku"], "sku", True), "quantity": _count(row["quantity"], "quantity")}
    for name in FIELDS:
        if name in result or name not in row:
            continue
        result[name] = (_price(row[name], name) if name in ("retail_price", "cost_price")
                        else _code(row[name] if row[name] is not None else "", optional=True) if name == "barcode"
                        else _text(row[name], name) if name in TEXT_LIMITS
                        else _count(row[name], name, name == "target_stock"))
    return result


def validate_rows(rows):
    if not isinstance(rows, (list, tuple)) or not rows or len(rows) > MAX_ROWS:
        raise InventoryError(f"Import must contain 1–{MAX_ROWS:,} rows.")
    normalized, seen, barcodes = [], {}, {}
    for index, row in enumerate(rows, 1):
        try:
            item = validate_row(row)
        except InventoryError as error:
            raise InventoryError(f"Row {index}: {error}") from None
        key = item["sku"].casefold()
        if key in seen:
            raise InventoryError(f"Row {index}: duplicate part number {item['sku']!r} (also row {seen[key]}).")
        seen[key] = index
        barcode = item.get("barcode", "")
        if barcode and barcode in barcodes:
            raise InventoryError(f"Row {index}: duplicate barcode {barcode!r} (also row {barcodes[barcode]}).")
        if barcode:
            barcodes[barcode] = index
        normalized.append(item)
    return normalized


def _cell_value(cell):
    value = cell.value
    # Excel stores numeric part codes without their displayed leading zeroes.
    # Honor only simple zero-padding formats; do not evaluate format expressions.
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 0 <= value < 10**15 and value == int(value):
        pattern = str(getattr(cell, "number_format", ""))
        if re.fullmatch(r"0{2,128}", pattern):
            return str(int(value)).zfill(len(pattern))
        if isinstance(value, float):
            return int(value)
    return value


def _read_spreadsheet(path, sheet=None, cancel=None):
    path = Path(path)
    if not path.is_file():
        raise InventoryError("Choose an existing spreadsheet file.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise InventoryError("Inventory files must be 50 MiB or smaller.")
    suffix = path.suffix.lower()
    _cancelled(cancel)
    if suffix in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise InventoryError("Excel import needs openpyxl. Install the desktop requirements.") from None
        try:
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                    raise InventoryError("The expanded Excel workbook is too large.")
            book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
            raise InventoryError("Cannot read this Excel workbook. Save a fresh .xlsx copy and try again.") from error
        try:
            sheets = list(book.sheetnames)
            if sheet is not None and sheet not in sheets:
                raise InventoryError("The selected worksheet does not exist.")
            selected = sheet or sheets[0]
            worksheet = book[selected]
            rows = []
            # Ignore untrusted sheet dimensions; cap actual iteration ourselves.
            worksheet.reset_dimensions()
            for index, cells in enumerate(worksheet.iter_rows(max_col=MAX_COLUMNS + 1)):
                if index > MAX_ROWS:
                    raise InventoryError(f"Import at most {MAX_ROWS:,} data rows at a time.")
                _cancelled(cancel)
                values = [_cell_value(cell) for cell in cells]
                if len(values) > MAX_COLUMNS and values[-1] not in (None, ""):
                    raise InventoryError(f"Import at most {MAX_COLUMNS} columns.")
                rows.append(values[:MAX_COLUMNS])
            return rows, sheets, selected
        finally:
            book.close()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError:
            raise InventoryError("Legacy Excel import needs xlrd. Install the desktop requirements.") from None
        try:
            book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
        except Exception as error:
            raise InventoryError("Cannot read this .xls workbook. Save a fresh Excel copy and try again.") from error
        try:
            sheets = book.sheet_names()
            if sheet is not None and sheet not in sheets:
                raise InventoryError("The selected worksheet does not exist.")
            selected = sheet or sheets[0]
            worksheet = book.sheet_by_name(selected)
            if worksheet.nrows > MAX_ROWS + 1 or worksheet.ncols > MAX_COLUMNS:
                raise InventoryError(f"Import at most {MAX_ROWS:,} rows and {MAX_COLUMNS} columns.")
            rows = []
            for index in range(worksheet.nrows):
                _cancelled(cancel)
                values = []
                for column in range(worksheet.ncols):
                    cell = worksheet.cell(index, column)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_NUMBER and 0 <= value < 10**15 and value == int(value):
                        fmt = book.format_map.get(book.xf_list[worksheet.cell_xf_index(index, column)].format_key)
                        pattern = fmt.format_str if fmt else ""
                        value = str(int(value)).zfill(len(pattern)) if re.fullmatch(r"0{2,128}", pattern) else int(value)
                    values.append(value)
                rows.append(values)
            return rows, sheets, selected
        finally:
            book.release_resources()
    if suffix == ".csv":
        raw = path.read_bytes()
        _cancelled(cancel)
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                content = raw.decode("utf-16") if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else raw.decode("cp1252")
            except UnicodeDecodeError:
                raise InventoryError("Save this CSV as UTF-8 and try again.") from None
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = []
        try:
            for index, row in enumerate(csv.reader(io.StringIO(content), dialect)):
                _cancelled(cancel)
                if index > MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise InventoryError(f"Import at most {MAX_ROWS:,} rows and {MAX_COLUMNS} columns.")
                rows.append(row)
        except csv.Error:
            raise InventoryError("Cannot read this CSV. Check its quoting and delimiters.") from None
        return rows, [], ""
    raise InventoryError("Choose an .xlsx, .xlsm, .xls, or .csv file.")


def _headers(rows):
    if not rows:
        raise InventoryError("The spreadsheet is empty.")
    while rows and all(value is None or str(value).strip() == "" for value in rows[-1]):
        rows.pop()
    if not rows:
        raise InventoryError("The spreadsheet is empty.")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise InventoryError("The first row must contain column headings.")
    return headers


def inspect_spreadsheet(path, sheet=None, cancel=None):
    rows, sheets, selected = _read_spreadsheet(path, sheet, cancel)
    headers = _headers(rows)
    return InventoryPreview(headers, [row[:len(headers)] for row in rows[1:11]], sheets,
                            max(0, len(rows) - 1), suggest_mapping(headers), Path(path).name, selected)


def parse_excel(path, mapping=None, sheet=None, cancel=None, progress=None, include_zero=True):
    rows, _sheets, selected = _read_spreadsheet(path, sheet, cancel)
    headers = _headers(rows)
    mapping = suggest_mapping(headers) if mapping is None else dict(mapping)
    resolved = {}
    errors, warnings, result, seen, barcodes = [], [], [], {}, {}
    skipped_zero = 0
    for name, column in mapping.items():
        if name not in FIELDS or column is None or column == "":
            continue
        if isinstance(column, str):
            matches = [index for index, heading in enumerate(headers) if heading == column]
            if len(matches) != 1:
                errors.append(f"Choose one unambiguous column for {name.replace('_', ' ')}.")
                continue
            column = matches[0]
        if isinstance(column, bool) or not isinstance(column, int) or not 0 <= column < len(headers):
            errors.append(f"Invalid column for {name.replace('_', ' ')}.")
        else:
            resolved[name] = column
    for name in ("sku", "quantity"):
        if name not in resolved:
            errors.append(f"Map a column to {name.replace('_', ' ')} before importing.")
    if errors:
        return InventoryImportPlan([], errors, warnings, Path(path).name, headers, resolved, selected)
    missing = [name.replace("_", " ") for name in ("description", "location", "reorder_point", "supplier") if name not in resolved]
    if missing:
        warnings.append("Columns not mapped: " + ", ".join(missing) + ". Existing values stay unchanged; new parts use blank text and a restock point of zero.")
    for index, values in enumerate(rows[1:], 2):
        _cancelled(cancel)
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        source = {name: values[column] if column < len(values) else None for name, column in resolved.items()}
        try:
            barcode = source.get("barcode")
            # Numeric Excel cells with a simple zero-padding format have already
            # become strings in _cell_value. Unformatted numbers are usable only
            # while within Excel's 15-digit integer precision; never guess a GTIN.
            if isinstance(barcode, (int, float)) and not isinstance(barcode, bool):
                if not math.isfinite(barcode) or barcode < 0 or barcode >= 10**15 or barcode != int(barcode):
                    raise InventoryError("Save this barcode as text in Excel; numeric cells can lose digits.")
                source["barcode"] = str(int(barcode))
            item = validate_row(source)
            if not include_zero and item["quantity"] == 0:
                skipped_zero += 1
                continue
            key = item["sku"].casefold()
            if key in seen:
                raise InventoryError(f"Duplicate part number {item['sku']!r}; also appears on row {seen[key]}.")
            seen[key] = index
            barcode = item.get("barcode", "")
            if barcode and barcode in barcodes:
                raise InventoryError(f"Duplicate barcode {barcode!r}; also appears on row {barcodes[barcode]}.")
            if barcode:
                barcodes[barcode] = index
            result.append(item)
        except InventoryError as error:
            if len(errors) < 100:
                errors.append(f"Row {index}: {error}")
        if progress is not None and index % 250 == 0:
            progress(index - 1, len(rows) - 1)
    if skipped_zero:
        warnings.append(f"Skipped {skipped_zero:,} zero-quantity rows. Enable Include zero-quantity parts to import them.")
    if not result and not errors:
        errors.append("No stock rows remain after excluding zero quantities." if skipped_zero else "The spreadsheet has no stock rows to import.")
    return InventoryImportPlan(result, errors, warnings, Path(path).name, headers, resolved, selected)


class InventoryStore:
    def __init__(self, path, *, clock=time.time):
        if str(path).startswith(("\\\\", "//")):
            raise InventoryError("Keep the inventory database on this PC's local disk.")
        self.path = Path(path).expanduser().resolve()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.clock = clock
        with closing(self._connect()) as connection, connection:
            connection.executescript("""
                CREATE TABLE IF NOT EXISTS inventory_state (
                    id INTEGER PRIMARY KEY CHECK(id=1), revision INTEGER NOT NULL DEFAULT 0,
                    last_import_at TEXT NOT NULL DEFAULT '', last_import_source TEXT NOT NULL DEFAULT ''
                );
                INSERT OR IGNORE INTO inventory_state(id) VALUES(1);
                CREATE TABLE IF NOT EXISTS inventory_items (
                    key TEXT PRIMARY KEY, sku TEXT NOT NULL, description TEXT NOT NULL DEFAULT '',
                    quantity INTEGER NOT NULL, location TEXT NOT NULL DEFAULT '',
                    reorder_point INTEGER NOT NULL DEFAULT 0, reorder_quantity INTEGER NOT NULL DEFAULT 0,
                    target_stock INTEGER, supplier TEXT NOT NULL DEFAULT '', revision INTEGER NOT NULL,
                    updated_at TEXT NOT NULL, source TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS inventory_description ON inventory_items(description COLLATE NOCASE);
                CREATE TABLE IF NOT EXISTS inventory_movements (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, sku_key TEXT NOT NULL, sku TEXT NOT NULL,
                    kind TEXT NOT NULL, quantity INTEGER NOT NULL, before_quantity INTEGER NOT NULL,
                    after_quantity INTEGER NOT NULL, delta INTEGER NOT NULL,
                    reason TEXT NOT NULL DEFAULT '', reference TEXT NOT NULL DEFAULT '',
                    actor TEXT NOT NULL DEFAULT '', actor_id INTEGER,
                    created_at TEXT NOT NULL, item_revision INTEGER NOT NULL,
                    changes_json TEXT NOT NULL DEFAULT '{}'
                );
                CREATE INDEX IF NOT EXISTS inventory_movement_part ON inventory_movements(sku_key,id DESC);
                CREATE TABLE IF NOT EXISTS inventory_movement_requests (
                    actor_key TEXT NOT NULL, request_id TEXT NOT NULL, payload_json TEXT NOT NULL,
                    result_json TEXT NOT NULL, PRIMARY KEY(actor_key,request_id)
                );
            """)
            connection.execute("BEGIN IMMEDIATE")
            columns = {row[1] for row in connection.execute("PRAGMA table_info(inventory_items)")}
            if "barcode" not in columns:
                connection.execute("ALTER TABLE inventory_items ADD COLUMN barcode TEXT NOT NULL DEFAULT ''")
            for name, default in (("alternate_part_number", ""), ("substitute_part_number", ""),
                                  ("retail_price", "0.00"), ("cost_price", "0.00")):
                if name not in columns:
                    connection.execute(f"ALTER TABLE inventory_items ADD COLUMN {name} TEXT NOT NULL DEFAULT '{default}'")
            connection.execute("CREATE UNIQUE INDEX IF NOT EXISTS inventory_barcode ON inventory_items(barcode) WHERE barcode<>''")

    def _connect(self):
        connection = sqlite3.connect(self.path, timeout=5)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA busy_timeout=5000")
        return connection

    @staticmethod
    def _public(row):
        if row is None:
            return None
        item = dict(row)
        item.pop("key", None)
        item["low_stock"] = item["quantity"] <= item["reorder_point"]
        target = item["target_stock"] if item["target_stock"] is not None else item["reorder_point"] + 1
        item["suggested_restock"] = max(item["reorder_quantity"], target - item["quantity"], 0) if item["low_stock"] else 0
        return item

    def get(self, sku):
        with closing(self._connect()) as connection:
            return self._public(connection.execute("SELECT * FROM inventory_items WHERE key=?", (str(sku).strip().casefold(),)).fetchone())

    def resolve_code(self, code):
        """Resolve a scanner's full text without changing stock or its revision."""
        code = _code(code)
        with closing(self._connect()) as connection:
            # Separate indexed equality probes avoid a table scan or LIKE
            # wildcard interpretation. A code matching two parts is not safe.
            rows = connection.execute("""SELECT * FROM inventory_items WHERE key=?
                UNION SELECT * FROM inventory_items WHERE barcode=? AND barcode<>''""",
                                      (code.casefold(), code)).fetchall()
        if not rows:
            raise InventoryError("No part matches this barcode or exact part number. Check the label or import its barcode.", 404)
        if len(rows) != 1:
            raise InventoryError("This code matches different parts as a barcode and a part number. Correct the catalogue before scanning it.", 409)
        return self._public(rows[0])

    def search(self, query="", limit=100, offset=0, low_stock=False):
        if not isinstance(query, str) or len(query) > 500:
            raise InventoryError("Search must contain at most 500 characters.")
        limit = max(1, min(int(limit), 500))
        offset = max(0, int(offset))
        where, parameters = [], []
        for word in query.split():
            pattern = "%" + word.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
            where.append("(sku LIKE ? ESCAPE '\\' OR description LIKE ? ESCAPE '\\' OR location LIKE ? ESCAPE '\\' OR supplier LIKE ? ESCAPE '\\' OR barcode LIKE ? ESCAPE '\\' OR alternate_part_number LIKE ? ESCAPE '\\' OR substitute_part_number LIKE ? ESCAPE '\\')")
            parameters.extend([pattern] * 7)
        if low_stock:
            where.append("quantity <= reorder_point")
        sql = "SELECT * FROM inventory_items" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY sku COLLATE NOCASE LIMIT ? OFFSET ?"
        with closing(self._connect()) as connection:
            return [self._public(row) for row in connection.execute(sql, (*parameters, limit, offset))]

    def low_stock(self, limit=100, offset=0):
        return self.search(limit=limit, offset=offset, low_stock=True)

    def summary(self):
        with closing(self._connect()) as connection:
            connection.execute("BEGIN")
            return self._summary(connection)

    def preview_import(self, plan):
        if isinstance(plan, InventoryImportPlan):
            if not plan.valid:
                raise InventoryError("Resolve spreadsheet errors before importing.")
            rows = plan.rows
        else:
            rows = plan
        normalized = validate_rows(rows)
        created = updated = unchanged = 0
        # Report changed fields, including replacements of existing quantities.
        changes = []
        with closing(self._connect()) as connection:
            connection.execute("BEGIN")
            self._validate_barcodes(connection, normalized)
            summary = self._summary(connection)
            for item in normalized:
                old = connection.execute("SELECT * FROM inventory_items WHERE key=?", (item["sku"].casefold(),)).fetchone()
                if old is None:
                    created += 1
                    if len(changes) < 100:
                        changes.append({"sku": item["sku"], "kind": "new", "before": None, "after": item})
                else:
                    fields = [name for name, value in item.items() if old[name] != value]
                    if fields:
                        updated += 1
                        if len(changes) < 100:
                            changes.append({"sku": item["sku"], "kind": "update", "fields": fields,
                                            "before": {name: old[name] for name in fields}, "after": {name: item[name] for name in fields}})
                    else:
                        unchanged += 1
        return {"revision": summary["revision"], "created": created, "updated": updated,
                "unchanged": unchanged, "rows": len(normalized), "changes": changes}

    def close(self):
        """Connections are short lived; provided for service lifecycle symmetry."""

    @staticmethod
    def _summary(connection):
        state = dict(connection.execute("SELECT revision,last_import_at,last_import_source FROM inventory_state WHERE id=1").fetchone())
        totals = connection.execute("SELECT COUNT(*),COALESCE(SUM(quantity),0),COALESCE(SUM(quantity<=reorder_point),0) FROM inventory_items").fetchone()
        return {**state, "item_count": totals[0], "total_quantity": totals[1], "low_stock_count": totals[2]}

    @staticmethod
    def _validate_barcodes(connection, normalized):
        """Check the final assignment, allowing deliberate barcode swaps."""
        owners = {row["barcode"]: row["key"] for row in connection.execute(
            "SELECT key,barcode FROM inventory_items WHERE barcode<>''")}
        changed_keys = {row["sku"].casefold() for row in normalized if "barcode" in row}
        owners = {code: key for code, key in owners.items() if key not in changed_keys}
        for item in normalized:
            code = item.get("barcode", "")
            if not code:
                continue
            key = item["sku"].casefold()
            if code in owners and owners[code] != key:
                raise InventoryError(f"Barcode {code!r} is assigned to another part. Each barcode must identify one part.", 409)
            owners[code] = key

    @staticmethod
    def _audit_text(value, label, limit, required=False):
        if not isinstance(value, str):
            raise InventoryError(f"{label} must contain text.")
        value = value.strip()
        if len(value) > limit or (required and not value) or any(ord(c) < 32 and c not in "\n\t" for c in value):
            raise InventoryError(f"{label} must contain {'1–' if required else 'at most '}{limit} readable characters.")
        return value

    @staticmethod
    def _actor_label(connection, actor, actor_id):
        if actor_id is None:
            return actor
        user = connection.execute("SELECT username,deleted_at FROM users WHERE id=?", (actor_id,)).fetchone()
        return user["username"] if user and user["deleted_at"] is None else "Deleted employee"

    @classmethod
    def _public_movement(cls, connection, row):
        movement = dict(row)
        movement.pop("sku_key", None)
        actor_id = movement.pop("actor_id", None)
        movement["actor"] = cls._actor_label(connection, movement["actor"], actor_id)
        movement["changes"] = json.loads(movement.pop("changes_json", "{}"))
        return movement

    def _insert_movement(self, connection, *, sku, kind, quantity, before, after,
                         reason, reference, actor, actor_id, timestamp, revision, changes=None):
        # Store server identities by ID only. Account deletion then anonymizes
        # history without preserving a second copy of the employee's username.
        stored_actor = "" if actor_id is not None else actor
        cursor = connection.execute("""INSERT INTO inventory_movements
            (sku_key,sku,kind,quantity,before_quantity,after_quantity,delta,reason,reference,
             actor,actor_id,created_at,item_revision,changes_json)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (sku.casefold(), sku, kind, quantity, before, after, after - before, reason,
             reference, stored_actor, actor_id, timestamp, revision,
             json.dumps(changes or {}, ensure_ascii=False, separators=(",", ":"))))
        return self._public_movement(connection, connection.execute(
            "SELECT * FROM inventory_movements WHERE id=?", (cursor.lastrowid,)).fetchone())

    def movements(self, sku="", limit=50, before_id=None):
        """Bounded keyset pagination; lookups and cache refreshes create no audit."""
        if not isinstance(sku, str) or len(sku) > 128:
            raise InventoryError("Choose a valid part number for stock history.")
        if type(limit) is not int or not 1 <= limit <= 200:
            raise InventoryError("History page size must be between 1 and 200.")
        if before_id is not None and (type(before_id) is not int or not 1 <= before_id < 2**63):
            raise InventoryError("History cursor must be a positive whole number.")
        where, values = [], []
        if sku.strip():
            where.append("sku_key=?")
            values.append(sku.strip().casefold())
        if before_id is not None:
            where.append("id<?")
            values.append(before_id)
        sql = "SELECT * FROM inventory_movements" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY id DESC LIMIT ?"
        with closing(self._connect()) as connection:
            connection.execute("BEGIN")
            rows = connection.execute(sql, (*values, limit + 1)).fetchall()
            movements, response_bytes = [], 0
            for row in rows[:limit]:
                movement = self._public_movement(connection, row)
                size = len(json.dumps(movement, ensure_ascii=False, separators=(",", ":")).encode("utf-8")) + 1
                # Imported before/after descriptions may each contain thousands
                # of Unicode characters. Bound actual bytes as well as row count
                # so every page fits the desktop transport's 1 MiB response cap.
                if movements and response_bytes + size > MAX_HISTORY_BYTES:
                    break
                movements.append(movement)
                response_bytes += size
        has_more = len(rows) > len(movements)
        return {"movements": movements, "has_more": has_more,
                "next_before_id": movements[-1]["id"] if has_more else None}

    def _record_movement(self, connection, sku, kind, quantity, *, expected_revision,
                         request_id, reason="", reference="", actor="Local user", actor_id=None):
        sku = _text(sku, "sku", True)
        if kind not in ("receive", "sale", "return", "count"):
            raise InventoryError("Choose receive, sale, return, or count for a stock movement.")
        if type(quantity) is not int or not (0 if kind == "count" else 1) <= quantity <= COUNT_LIMIT:
            raise InventoryError("Enter a positive whole quantity; a physical count can be zero.")
        if type(expected_revision) is not int or not 0 <= expected_revision < 2**63 - 1:
            raise InventoryError("Refresh this part before recording its stock movement.")
        if not isinstance(request_id, str) or not re.fullmatch(r"[0-9a-f]{32}", request_id):
            raise InventoryError("Stock movements need a unique 32-character request ID.")
        reason = self._audit_text(reason, "Reason", 500, kind == "count")
        reference = self._audit_text(reference, "Reference", 200)
        actor = self._audit_text(actor, "Employee name", 128, True)
        actor_key = "user:" + str(actor_id) if actor_id is not None else "local:" + actor
        payload = {"sku": sku.casefold(), "kind": kind, "quantity": quantity,
                   "expected_revision": expected_revision, "reason": reason, "reference": reference}
        encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        previous = connection.execute("SELECT payload_json,result_json FROM inventory_movement_requests WHERE actor_key=? AND request_id=?",
                                      (actor_key, request_id)).fetchone()
        if previous is not None:
            if previous["payload_json"] != encoded:
                raise InventoryError("This request ID was already used for a different stock movement. Refresh before trying again.", 409)
            result = json.loads(previous["result_json"])
            result["movement"]["actor"] = self._actor_label(connection, actor, actor_id)
            return result
        old = connection.execute("SELECT * FROM inventory_items WHERE key=?", (sku.casefold(),)).fetchone()
        if old is None:
            raise InventoryError("Part number was not found. Scan or select an existing part first.", 404)
        if old["revision"] != expected_revision:
            raise InventoryError("This part changed on another PC. Scan or refresh it and review the new quantity.", 409)
        before = old["quantity"]
        after = quantity if kind == "count" else before + (-quantity if kind == "sale" else quantity)
        if not 0 <= after <= COUNT_LIMIT:
            raise InventoryError("This movement would make stock negative or exceed the maximum quantity.", 409)
        revision = connection.execute("SELECT revision FROM inventory_state WHERE id=1").fetchone()[0] + 1
        timestamp = datetime.fromtimestamp(self.clock(), timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
        connection.execute("UPDATE inventory_items SET quantity=?,revision=?,updated_at=?,source=? WHERE key=?",
                           (after, revision, timestamp, "Stock movement", sku.casefold()))
        connection.execute("UPDATE inventory_state SET revision=? WHERE id=1", (revision,))
        movement = self._insert_movement(connection, sku=old["sku"], kind=kind, quantity=quantity,
            before=before, after=after, reason=reason, reference=reference, actor=actor, actor_id=actor_id,
            timestamp=timestamp, revision=revision, changes={"quantity": {"before": before, "after": after}})
        result = {"item": self._public(connection.execute("SELECT * FROM inventory_items WHERE key=?", (sku.casefold(),)).fetchone()),
                  "movement": movement, "summary": self._summary(connection)}
        stored_result = {**result, "movement": {**movement, "actor": "" if actor_id is not None else actor}}
        connection.execute("INSERT INTO inventory_movement_requests VALUES(?,?,?,?)",
                           (actor_key, request_id, encoded, json.dumps(stored_result, ensure_ascii=False)))
        return result

    def record_movement(self, sku, kind, quantity, *, expected_revision, request_id,
                        reason="", reference="", actor="Local user"):
        """Commit the quantity, audit, and retry receipt under one write lock."""
        with closing(self._connect()) as connection, connection:
            connection.execute("BEGIN IMMEDIATE")
            return self._record_movement(connection, sku, kind, quantity,
                expected_revision=expected_revision, request_id=request_id, reason=reason,
                reference=reference, actor=actor)

    def _apply_rows(self, connection, normalized, source, expected_revision, *, actor="Local user", actor_id=None, audit=True):
        """Apply already validated rows in the caller's write transaction."""
        timestamp = datetime.fromtimestamp(self.clock(), timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
        revision = connection.execute("SELECT revision FROM inventory_state WHERE id=1").fetchone()[0]
        if expected_revision is not None and expected_revision != revision:
            raise InventoryError("Stock changed since the preview. Refresh and review this import again.", 409)
        self._validate_barcodes(connection, normalized)
        # Remember the originals before releasing changed barcode assignments.
        # Releasing them together permits an explicit A/B barcode swap without
        # violating the uniqueness constraint halfway through the transaction.
        originals = {row["sku"].casefold(): connection.execute(
            "SELECT * FROM inventory_items WHERE key=?", (row["sku"].casefold(),)).fetchone()
            for row in normalized}
        existing_count = connection.execute("SELECT COUNT(*) FROM inventory_items").fetchone()[0]
        incoming_new = sum(old is None for old in originals.values())
        if existing_count + incoming_new > MAX_ROWS:
            raise InventoryError(f"The inventory can contain at most {MAX_ROWS:,} different parts.", 413)
        connection.executemany("UPDATE inventory_items SET barcode='' WHERE key=?",
            ((row["sku"].casefold(),) for row in normalized if "barcode" in row
             and originals[row["sku"].casefold()] is not None
             and originals[row["sku"].casefold()]["barcode"] != row["barcode"]))
        revision += 1
        created = updated = unchanged = 0
        for item in normalized:
            key = item["sku"].casefold()
            old = originals[key]
            changes = {name: {"before": old[name] if old is not None else None, "after": value}
                       for name, value in item.items() if old is None or old[name] != value}
            if old is None:
                created += 1
                values = {"description": "", "location": "", "reorder_point": 0,
                          "reorder_quantity": 0, "target_stock": None, "supplier": "", "barcode": "",
                          "alternate_part_number": "", "substitute_part_number": "",
                          "retail_price": "0.00", "cost_price": "0.00", **item}
                names = ("key", *FIELDS, "revision", "updated_at", "source")
                connection.execute(f"INSERT INTO inventory_items({','.join(names)}) VALUES({','.join('?' for _ in names)})",
                                   (key, *(values[name] for name in FIELDS), revision, timestamp, source))
            elif changes:
                updated += 1
                assignments = ",".join(f"{name}=?" for name in item)
                connection.execute(f"UPDATE inventory_items SET {assignments},revision=?,updated_at=?,source=? WHERE key=?",
                                   (*item.values(), revision, timestamp, source, key))
            else:
                unchanged += 1
            if audit and changes:
                before = old["quantity"] if old is not None else 0
                self._insert_movement(connection, sku=item["sku"], kind="import", quantity=item["quantity"],
                    before=before, after=item["quantity"], reason="Stock imported or edited", reference=source,
                    actor=actor, actor_id=actor_id, timestamp=timestamp, revision=revision, changes=changes)
        connection.execute("UPDATE inventory_state SET revision=?,last_import_at=?,last_import_source=? WHERE id=1", (revision, timestamp, source))
        return {**self._summary(connection), "created": created, "updated": updated, "unchanged": unchanged, "rows": len(normalized)}

    def import_rows(self, rows, source="", expected_revision=None):
        normalized = validate_rows(rows)
        source = str(source).strip()
        if len(source) > 255 or any(ord(c) < 32 for c in source):
            raise InventoryError("Import source must be a readable name of at most 255 characters.")
        if expected_revision is not None and (isinstance(expected_revision, bool) or not isinstance(expected_revision, int) or expected_revision < 0):
            raise InventoryError("Expected revision must be a nonnegative whole number.")
        with closing(self._connect()) as connection, connection:
            connection.execute("BEGIN IMMEDIATE")
            return self._apply_rows(connection, normalized, source, expected_revision)

    def replace_snapshot(self, rows, revision, last_import_at="", last_import_source=""):
        """Replace an account-scoped local cache with an authoritative server snapshot."""
        if not isinstance(rows, (list, tuple)) or len(rows) > MAX_ROWS:
            raise InventoryError("The stock snapshot is too large.")
        normalized = validate_rows([{name: row[name] for name in FIELDS if name in row} for row in rows]) if rows else []
        if isinstance(revision, bool) or not isinstance(revision, int) or revision < 0:
            raise InventoryError("Invalid inventory snapshot revision.")
        with closing(self._connect()) as connection, connection:
            connection.execute("BEGIN IMMEDIATE")
            connection.execute("DELETE FROM inventory_items")
            if normalized:
                self._apply_rows(connection, normalized, str(last_import_source)[:255], None, audit=False)
            connection.execute("UPDATE inventory_items SET revision=?", (revision,))
            for row in rows:
                row_revision = row.get("revision", revision)
                updated_at = row.get("updated_at", last_import_at)
                source = row.get("source", last_import_source)
                if isinstance(row_revision, bool) or not isinstance(row_revision, int) or not 0 <= row_revision <= revision:
                    raise InventoryError("Invalid cached item revision.")
                if not isinstance(updated_at, str) or len(updated_at) > 80 or not isinstance(source, str) or len(source) > 255:
                    raise InventoryError("Invalid cached item source metadata.")
                connection.execute("UPDATE inventory_items SET revision=?,updated_at=?,source=? WHERE key=?",
                                   (row_revision, updated_at, source, str(row["sku"]).strip().casefold()))
            connection.execute("UPDATE inventory_state SET revision=?,last_import_at=?,last_import_source=? WHERE id=1",
                               (revision, str(last_import_at)[:80], str(last_import_source)[:255]))
        return self.summary()
